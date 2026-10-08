import copy
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from seculens import scan_sbom, write_word_report
from seculens.cli import main
from seculens.customer_report import ordered_findings
from seculens.severity import cvss3_base, record_severity

ROOT = Path(__file__).resolve().parents[1]
VECTORS = [
    ("CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H", 9.8),
    ("CVSS:3.0/AV:N/AC:L/PR:N/UI:N/S:C/C:H/I:H/A:H", 10.0),
    ("CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:C/C:H/I:H/A:H", 9.9),
    ("CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:N/I:N/A:L", 5.3),
    ("CVSS:3.1/AV:N/AC:L/PR:N/UI:R/S:C/C:L/I:L/A:N", 6.1),
    ("CVSS:3.1/AV:N/AC:L/PR:H/UI:N/S:U/C:H/I:H/A:H", 7.2),
    ("CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:N/I:N/A:N", 0.0),
]


@pytest.mark.parametrize("vector,score", VECTORS)
def test_cvss_base_reference_values(vector, score):
    assert cvss3_base(vector) == score
    assert cvss3_base(vector + "/E:U/RL:O/RC:R") == score  # Base, not temporal score.


@pytest.mark.parametrize(
    "vector",
    [
        None,
        "9.8",
        "CVSS:4.0/AV:N",
        "CVSS:3.1/AV:N",
        VECTORS[0][0] + "/AV:L",
        VECTORS[0][0].replace("AC:L", "AC:INVALID"),
        VECTORS[0][0] + "/E:INVALID",
    ],
)
def test_bad_or_unsupported_cvss_remains_unrated(vector):
    assert cvss3_base(vector) is None


def records():
    base = json.loads((ROOT / "examples/database.json").read_text())
    if isinstance(base, dict):
        base = base["records"]
    demo = base[0]
    return [
        dict(copy.deepcopy(demo), id="TEST-HIGH", database_specific={"severity": "HIGH"}),
        dict(
            copy.deepcopy(demo),
            id="TEST-CRITICAL",
            severity=[{"type": "CVSS_V3", "score": VECTORS[0][0]}],
        ),
        dict(
            copy.deepcopy(demo),
            id="TEST-UNRATED",
            severity=[{"type": "CVSS_V4", "score": "CVSS:4.0/AV:N"}],
        ),
    ]


def test_affected_override_and_unrelated_packages():
    advisory = records()[1]
    component = {"name": "lodash", "version": "4.17.20", "ecosystem": "npm"}
    advisory["affected"][0]["severity"] = [{"type": "CVSS_V3", "score": VECTORS[3][0]}]
    unrelated = copy.deepcopy(advisory["affected"][0])
    unrelated["package"]["name"] = "unrelated"
    unrelated["severity"][0]["score"] = VECTORS[0][0]
    advisory["affected"].append(unrelated)
    result = record_severity(advisory, component)
    assert result["level"] == "medium"
    assert result["score"] == 5.3
    assert result["sources"][0]["field"] == "affected[0].severity[0]"


def test_alias_merge_keeps_max_severity_and_all_sources():
    advisory = records()
    advisory[0]["aliases"] = ["TEST-CRITICAL"]
    report = scan_sbom((ROOT / "examples/cyclonedx.json").read_text(), advisory)
    findings = [f for f in report["findings"] if f["category"] == "vulnerability"]
    assert len(findings) == 2
    merged = next(f for f in findings if f["severity"]["level"] == "critical")
    assert merged["severity"]["score"] == 9.8
    assert {s["recordId"] for s in merged["severity"]["sources"]} == {"TEST-HIGH", "TEST-CRITICAL"}


@pytest.mark.parametrize("language", ["ja", "en"])
def test_customer_cover_tables_sort_colors_and_standard_compatibility(tmp_path, language):
    report = scan_sbom(
        (ROOT / "examples/cyclonedx.json").read_text(),
        records(),
        policy={"allow": ["MIT", "Apache-2.0"]},
        customer="Sample Customer",
    )
    before = copy.deepcopy(report)
    ordered = ordered_findings(report)
    assert [f["severity"]["level"] for f in ordered] == ["critical", "high", "unknown"]
    path = tmp_path / "customer.docx"
    write_word_report(report, path, language, style="customer", issuer="Example Security Team")
    doc = Document(path)
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    assert "Example Security Team" in paragraphs
    assert "F-001" in paragraphs
    assert len(doc.tables) == 4
    assert [row.cells[1].text for row in doc.tables[2].rows[1:]] == (
        ["重大", "高", "未評価"] if language == "ja" else ["Critical", "High", "Unrated"]
    )
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
        assert 'w:fill="FEE2E2"' in xml and 'w:color w:val="991B1B"' in xml
        assert "tblHeader" in xml and "cantSplit" in xml
        assert "titlePg" in xml
    assert report == before  # Presentation cannot reorder/mutate assessment JSON.
    standard = tmp_path / "standard.docx"
    write_word_report(report, standard, language)
    assert "Example Security Team" not in "\n".join(p.text for p in Document(standard).paragraphs)


def test_empty_customer_report_does_not_claim_safe(tmp_path):
    sbom = json.dumps({"bomFormat": "CycloneDX", "specVersion": "1.6", "components": []})
    report = scan_sbom(sbom, [])
    path = tmp_path / "empty.docx"
    write_word_report(report, path, "en", style="customer")
    paragraphs = "\n".join(p.text for p in Document(path).paragraphs)
    assert "No findings were recorded" in paragraphs
    assert "not absence of vulnerabilities" in paragraphs
    with pytest.raises(ValueError):
        write_word_report(report, path, style="invalid")


def test_cli_customer_mode_and_issuer(tmp_path):
    assert (
        main(
            [
                "scan",
                str(ROOT / "examples/cyclonedx.json"),
                "--db",
                str(ROOT / "examples/database.json"),
                "--report-style",
                "customer",
                "--issuer",
                "Example Team",
                "-o",
                str(tmp_path),
            ]
        )
        == 0
    )
    assert "Example Team" in "\n".join(
        p.text for p in Document(tmp_path / "report.docx").paragraphs
    )


def test_mixed_affected_entries_use_each_entry_effective_severity():
    advisory = records()[1]
    first = advisory["affected"][0]
    first["severity"] = [{"type": "CVSS_V3", "score": VECTORS[3][0]}]
    second = copy.deepcopy(first)
    second.pop("severity")
    advisory["affected"].append(second)
    result = record_severity(advisory, {"name": "lodash", "version": "4.17.20", "ecosystem": "npm"})
    assert result["level"] == "critical"
    assert {s["field"] for s in result["sources"]} == {"severity[0]", "affected[0].severity[0]"}


def test_report_id_changes_with_assessment_timestamp_and_evidence():
    from seculens.customer_report import report_id

    report = scan_sbom((ROOT / "examples/cyclonedx.json").read_text(), [])
    original = report_id(report)
    report["createdAt"] = "2026-10-09T01:00:00Z"
    assert report_id(report) != original
