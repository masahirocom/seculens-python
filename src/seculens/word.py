"""Customer Word reports with packaged Japanese font embedding."""

from __future__ import annotations

import uuid
from importlib.resources import files
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

_CATEGORIES = {
    "vulnerability": "既知の脆弱性",
    "license": "ライセンス評価",
    "security": "コードのセキュリティ確認",
    "quality": "コード品質",
    "coverage": "未判定と検査範囲",
}
_STATUS = {
    "affected": "脆弱性該当",
    "matched": "脆弱性該当",
    "no-match": "照合なし",
    "unassessed": "未判定",
    "review": "要確認",
    "denied": "ポリシー違反",
}
_TRANSLATIONS = {
    "Review the advisory references for a fixed version and validate the upgrade.": "参照先で修正済みバージョンを確認し、更新後の動作を検証してください。",
    "Review the call and trace untrusted input before deciding whether it is exploitable.": "処理を確認し、外部入力の経路を追跡して、実際に悪用可能か判断してください。",
    "Consider simplifying the function and adding focused tests.": "関数の分割や簡素化を検討し、分岐に対応したテストを追加してください。",
    "Confirm the license, usage and distribution conditions against the customer policy.": "ライセンスと利用・配布条件を確認し、顧客のポリシーに照らして判断してください。",
    "No match means no matching record in the supplied database snapshot, not absence of vulnerabilities.": "照合なしは使用したDB内に該当情報がないことを示し、脆弱性が存在しないことを保証しません。",
    "License policy checks are not a legal compliance determination.": "ライセンスポリシー評価は法的な適合性の確定を行うものではありません。",
    "AST rules identify review candidates and complexity, not proven exploitability.": "AST検査は要確認のコードや複雑性を抽出します。悪用可能性を立証するものではありません。",
    "SBOM content and package identifiers are supplied by the generating tool.": "構成部品の一覧と識別情報は、SBOM生成ツールから取得した情報に依存します。",
    "Packagist matching supports stable numeric versions and SemVer prereleases only; Composer branches and other version forms may be unassessed.": "Composerのブランチ指定など、未対応のバージョン表記は未判定になります。",
}


def _embed_font(document: Any) -> None:
    # ECMA-376 obfuscation: reversed GUID bytes repeated over the first 32 bytes.
    key = uuid.uuid4()
    mask = key.bytes[::-1]
    data = bytearray(files("seculens").joinpath("assets/NotoSansJP-Regular.otf").read_bytes())
    for index in range(32):
        data[index] ^= mask[index % 16]
    table = document.part.part_related_by(RT.FONT_TABLE)
    font = Part(
        PackURI("/word/fonts/seculens.odttf"),
        "application/vnd.openxmlformats-officedocument.obfuscatedFont",
        bytes(data),
        document.part.package,
    )
    relationship = table.relate_to(
        font, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
    )
    tree = etree.fromstring(table.blob)
    entry = OxmlElement("w:font")
    entry.set(qn("w:name"), "Noto Sans JP")
    embedded = OxmlElement("w:embedRegular")
    embedded.set(qn("r:id"), relationship)
    embedded.set(qn("w:fontKey"), "{" + str(key).upper() + "}")
    embedded.set(qn("w:subsetted"), "false")
    entry.append(embedded)
    tree.append(entry)
    table._blob = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    setting = OxmlElement("w:embedTrueTypeFonts")
    document.settings.element.append(setting)


def _reference(document: Any, url: str, text: str) -> None:
    paragraph = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), document.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


def write_word_report(report: dict[str, Any], file: str | Path, language: str = "en") -> None:
    if language not in ("en", "ja"):
        raise ValueError("Language must be en or ja")
    ja = language == "ja"
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(20)
    for name, size in (("Normal", 10.5), ("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = document.styles[name]
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
        style.font.name = "Noto Sans JP" if ja else "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), style.font.name)
        style.paragraph_format.space_after = Pt(6)
        if name != "Normal":
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    if ja:
        _embed_font(document)
    document.core_properties.author = "SecuLens"
    document.core_properties.title = "Software Security Assessment Report"
    p = document.add_paragraph

    def localize(text: str) -> str:
        return _TRANSLATIONS.get(text, text) if ja else text

    def status(value: str) -> str:
        return _STATUS.get(value, value) if ja else value

    findings = report["findings"]
    components = report["sbom"]["components"]
    vulnerable = len({f["subject"] for f in findings if f["category"] == "vulnerability"})
    unresolved = sum(c["status"] == "unassessed" for c in report["checks"])
    document.add_heading(
        "ソフトウェアセキュリティ検査報告書" if ja else "Software Security Assessment Report", 0
    )
    p(f"{report['customer']} | SecuLens {report['tool']['version']}")
    p(f"{'対象' if ja else 'Target'} {report['target']} | {report['createdAt']}")
    p(
        f"SBOM内の{len(components)}件の構成部品を評価しました。既知の脆弱性に該当する部品は{vulnerable}件、脆弱性照合が未判定の部品は{unresolved}件です。以下の指摘と確認事項に沿って対応してください。"
        if ja
        else f"Assessed {len(components)} SBOM components. {vulnerable} components matched known vulnerability records; {unresolved} components have incomplete vulnerability assessments. Review the findings and follow-up actions below."
    )
    document.add_heading("検査範囲と根拠" if ja else "Scope and Evidence", 1)
    p(f"{report['sbom']['format']} {report['sbom']['version']}")
    p(f"SBOM SHA256 {report['sbomSha256']}")
    p(f"DB {report['database']['source']}")
    p(f"DB SHA256 {report['database']['sha256']}")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "分類" if ja else "Category"
    table.rows[0].cells[1].text = "指摘件数" if ja else "Findings"
    for category, label in _CATEGORIES.items():
        cells = table.add_row().cells
        cells[0].text = label if ja else category
        cells[1].text = str(sum(f["category"] == category for f in findings))
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Mm(120), Mm(50)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement("w:" + edge)
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            item.set(qn("w:" + attr), value)
        borders.append(item)
    table._tbl.tblPr.append(borders)
    for row_index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.width = Mm(120 if column == 0 else 50)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7EDF3")
                cell._tc.get_or_add_tcPr().append(shading)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    for category, label in _CATEGORIES.items():
        group = [f for f in findings if f["category"] == category]
        if not group:
            continue
        document.add_heading(label if ja else category, 1)
        for finding in group:
            document.add_heading(f"{finding['ruleId']} {finding['summary']}", 2)
            component = next((c for c in components if c["id"] == finding["subject"]), None)
            subject = (
                f"{component['name']}@{component.get('version', 'unknown')}"
                if component
                else finding["subject"]
            )
            location = f" | line {finding['line']}" if finding.get("line") else ""
            p(f"{subject} | {status(finding['status'])}{location}")
            p(f"{'根拠' if ja else 'Evidence'} {finding['evidence']}")
            p(f"{'推奨対応' if ja else 'Action'} {localize(finding['recommendation'])}")
            if finding.get("aliases"):
                p(f"{'関連ID' if ja else 'Aliases'} {', '.join(finding['aliases'])}")
            references = sorted(
                set(finding.get("references", [])),
                key=lambda u: (0 if "/advisories/" in u else 1 if "nvd.nist.gov" in u else 2, u),
            )
            for index, reference in enumerate(references[:3], 1):
                url = urlsplit(reference)
                if url.scheme in ("http", "https") and url.netloc:
                    _reference(
                        document,
                        reference,
                        f"{'参照' if ja else 'Reference'} {index} {url.hostname}",
                    )
            if len(references) > 3:
                p(
                    "参照先の全一覧は同梱のJSON報告書に記録しています。"
                    if ja
                    else "The complete reference list is preserved in the accompanying JSON report."
                )
    document.add_heading("構成部品の評価状況" if ja else "Component Assessment Coverage", 1)
    for component in components:
        check = next((c for c in report["checks"] if c["componentId"] == component["id"]), {})
        p(
            f"{component['name']}@{component.get('version', 'unknown')} | {status(check.get('status', 'unassessed'))} | {'; '.join(component['licenses']) or 'License unknown'}"
        )
    document.add_heading("制約と確認事項" if ja else "Limitations and Review Notes", 1)
    for text in report["limitations"] + report["sbom"]["warnings"]:
        p(localize(text))
    document.save(file)
