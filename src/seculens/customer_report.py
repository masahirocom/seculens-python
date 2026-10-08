"""Formal customer report layout; colors supplement explicit severity labels."""

from __future__ import annotations

import hashlib
from typing import Any

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .severity import LEVELS

PALETTE = {
    "critical": ("重大", "Critical", "991B1B", "FEE2E2"),
    "high": ("高", "High", "B91C1C", "FFF1F2"),
    "medium": ("中", "Medium", "92400E", "FEF3C7"),
    "low": ("低", "Low", "075985", "E0F2FE"),
    "none": ("なし", "None", "166534", "F0FDF4"),
    "unknown": ("未評価", "Unrated", "475569", "F1F5F9"),
}
CATEGORIES = {
    "vulnerability": ("脆弱性", "Vulnerability"),
    "license": ("ライセンス", "License"),
    "security": ("コード確認", "Code review"),
    "quality": ("品質", "Quality"),
    "coverage": ("検査範囲", "Coverage"),
}


def report_id(report: dict[str, Any]) -> str:
    evidence = "|".join(
        report[key] if key != "database" else report[key]["sha256"]
        for key in ("sbomSha256", "database", "createdAt")
    )
    return "SL-" + hashlib.sha256(evidence.encode("utf-8")).hexdigest()[:12].upper()


def severity_level(finding: dict[str, Any]) -> str:
    value = finding.get("severity", {}).get("level", "unknown")
    return value if value in LEVELS else "unknown"


def ordered_findings(report: dict[str, Any]) -> list[dict[str, Any]]:
    return sorted(
        report["findings"],
        key=lambda f: (
            list(CATEGORIES).index(f["category"]),
            LEVELS.index(severity_level(f)) if f["category"] == "vulnerability" else 0,
            f["subject"],
            f["ruleId"],
            f.get("line", 0),
        ),
    )


def label(finding: dict[str, Any], ja: bool) -> str:
    if finding["category"] != "vulnerability":
        return (
            ("未判定" if ja else "Incomplete")
            if finding["status"] == "unassessed"
            else ("ポリシー違反" if ja else "Denied")
            if finding["status"] == "denied"
            else ("要確認" if ja else "Review")
        )
    return PALETTE[severity_level(finding)][0 if ja else 1]


def _shade(cell: Any, fill: str) -> None:
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(item)


def table(
    document: Any,
    rows: list[list[str]],
    widths: list[float],
    *,
    marks: dict[int, str] | None = None,
    compact: bool = False,
) -> Any:
    result = document.add_table(rows=0, cols=len(widths))
    result.autofit = False
    for column, width in zip(result.columns, widths, strict=True):
        column.width = Mm(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement("w:" + edge)
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            item.set(qn("w:" + attr), value)
        borders.append(item)
    result._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:w"), "80")
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    result._tbl.tblPr.append(margins)
    for index, values in enumerate(rows):
        row = result.add_row()
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for column, (cell, text, width) in enumerate(zip(row.cells, values, widths, strict=True)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            _shade(cell, "E7EDF3" if index == 0 else "FFFFFF" if index % 2 else "F8FAFC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if width <= 30 else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    run.font.size = Pt(9 if compact else 10)
                    run.bold = index == 0
            if index and column == 0 and marks and index in marks:
                color, fill = PALETTE[marks[index]][2:]
                _shade(cell, fill)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor.from_string(color)
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return result


def front(
    document: Any, report: dict[str, Any], findings: list[dict[str, Any]], ja: bool, issuer: str
) -> None:
    brand = document.add_paragraph()
    brand.paragraph_format.space_before = Pt(75)
    run = brand.add_run("SecuLens")
    run.bold, run.font.size = True, Pt(28)
    document.add_heading(
        "ソフトウェアセキュリティ評価報告書" if ja else "Software Security Assessment Report", 0
    )
    document.add_paragraph(
        "脆弱性とライセンスの評価およびコード確認"
        if ja
        else "Vulnerability and license assessment with code review"
    )
    customer = document.add_paragraph()
    customer.paragraph_format.space_before = Pt(36)
    customer.add_run(f"{report['customer']} {'御中' if ja else ''}".strip()).bold = True
    document.add_paragraph(f"{'検査対象' if ja else 'Target'}  {report['target']}")
    document.add_paragraph(f"{'検査日時' if ja else 'Assessment time'}  {report['createdAt']}")
    document.add_paragraph(f"{'報告書ID' if ja else 'Report ID'}  {report_id(report)}")
    if issuer:
        document.add_paragraph(f"{'作成者' if ja else 'Prepared by'}  {issuer}")
    document.add_paragraph(f"SecuLens {report['tool']['version']}")
    scope = document.add_paragraph(
        "本報告書は、対象の構成部品と脆弱性DBの照合、ライセンスポリシーの評価、および実施したコード検査の結果をまとめたものです。指摘一覧と詳細の根拠を参照し、対応方針を決定してください。"
        if ja
        else "This report summarizes the supplied component inventory, advisory matching, license policy evaluation and any code checks performed. Use the findings and supporting evidence to decide follow-up actions."
    )
    scope.paragraph_format.space_before = Pt(30)
    document.add_page_break()
    document.add_heading("1 検査結果の要約" if ja else "1 Assessment Summary", 1)
    vulnerabilities = [f for f in findings if f["category"] == "vulnerability"]
    incomplete = sum(c["status"] == "unassessed" for c in report["checks"])
    document.add_paragraph(
        f"構成部品{len(report['sbom']['components'])}件を照合し、{len(vulnerabilities)}件の脆弱性指摘を検出しました。該当する構成部品は{len({f['subject'] for f in vulnerabilities})}件、脆弱性照合が未判定の構成部品は{incomplete}件です。件数は同一部品内の別名IDを統合した後の値です。"
        if ja
        else f"Assessed {len(report['sbom']['components'])} components and identified {len(vulnerabilities)} vulnerability findings affecting {len({f['subject'] for f in vulnerabilities})} components. {incomplete} components have incomplete vulnerability assessments. Counts consolidate advisory aliases per component."
    )
    urgent = sum(severity_level(f) in ("critical", "high") for f in vulnerabilities)
    document.add_paragraph(
        f"重大または高の指摘は{urgent}件です。優先して参照先の修正情報と利用状況を確認してください。未評価の指摘も、重要度を確定するための確認が必要です。"
        if ja
        else f"{urgent} findings are Critical or High. Prioritize review of available fixes and deployment context. Unrated findings also require severity review."
    )
    actions = (
        [
            "優先して修正情報と影響を確認",
            "優先して修正情報と影響を確認",
            "更新計画と影響を確認",
            "通常の更新計画で確認",
            "評価根拠と利用条件を確認",
            "重要度の根拠を追加確認",
        ]
        if ja
        else [
            "Prioritize fixes and impact review",
            "Prioritize fixes and impact review",
            "Review upgrade plan and impact",
            "Review in routine upgrade planning",
            "Review evidence and usage context",
            "Obtain severity evidence",
        ]
    )
    rows = [
        [
            "重要度" if ja else "Severity",
            "件数" if ja else "Count",
            "確認方針" if ja else "Review guidance",
        ]
    ]
    rows += [
        [
            PALETTE[level][0 if ja else 1],
            str(sum(severity_level(f) == level for f in vulnerabilities)),
            actions[index],
        ]
        for index, level in enumerate(LEVELS)
    ]
    table(
        document,
        rows,
        [30, 20, 120],
        marks={index + 1: level for index, level in enumerate(LEVELS)},
    )
    document.add_paragraph(
        "重要度はCVSS 3.0／3.1の基本値、またはDB提供のラベルに基づきます。複数の根拠がある場合は最も高い重要度を採用し、詳細に根拠を残します。CVSS 2／4のベクトルは値を計算せず保存し、他に判定可能な根拠がなければ未評価とします。重要度は顧客環境での悪用可能性や事業影響を確定するものではありません。"
        if ja
        else "Severity uses CVSS 3.0/3.1 base scores or database labels. The highest available severity is displayed when sources differ; all evidence is retained. CVSS 2/4 vectors are preserved but not calculated; absent other usable evidence they remain Unrated. Severity does not establish exploitability or business impact in the customer environment."
    )
    rows = [["分類" if ja else "Category", "件数" if ja else "Count"]] + [
        [names[0 if ja else 1], str(sum(f["category"] == category for f in findings))]
        for category, names in CATEGORIES.items()
    ]
    table(document, rows, [140, 30], compact=True)
    document.add_page_break()
    document.add_heading("2 指摘事項の一覧" if ja else "2 Findings Register", 1)
    document.add_paragraph(
        "番号は後続の詳細と対応します。コード検査とライセンスの要確認事項は脆弱性の重要度と区別して記載しています。"
        if ja
        else "Finding numbers correspond to the details that follow. Code and license review candidates are distinguished from vulnerability severity."
    )
    rows = [
        [
            "番号" if ja else "No.",
            "重要度／状態" if ja else "Severity / status",
            "分類" if ja else "Category",
            "対象と指摘" if ja else "Subject and finding",
        ]
    ]
    for index, finding in enumerate(findings, 1):
        component = next(
            (c for c in report["sbom"]["components"] if c["id"] == finding["subject"]), None
        )
        subject = (
            f"{component['name']}@{component.get('version', 'unknown')}"
            if component
            else finding["subject"] + (f":{finding['line']}" if finding.get("line") else "")
        )
        rows.append(
            [
                f"F-{index:03d}",
                label(finding, ja),
                CATEGORIES[finding["category"]][0 if ja else 1],
                f"{subject}\n{finding['ruleId']}",
            ]
        )
    if findings:
        result = table(document, rows, [18, 28, 28, 96], compact=True)
        for index, finding in enumerate(findings, 1):
            if finding["category"] == "vulnerability":
                cell = result.rows[index].cells[1]
                color, fill = PALETTE[severity_level(finding)][2:]
                _shade(cell, fill)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(color)
    else:
        document.add_paragraph(
            "指摘事項はありません。検査範囲と制約もご確認ください。"
            if ja
            else "No findings were recorded. Review assessment coverage and limitations as well."
        )
    document.add_page_break()
    document.add_heading("3 指摘事項の詳細" if ja else "3 Finding Details", 1)


def severity_detail(document: Any, finding: dict[str, Any], ja: bool) -> None:
    if finding["category"] != "vulnerability":
        return
    level = severity_level(finding)
    severity = finding.get("severity", {})
    score = (
        f" | CVSS 3 基本値 {severity['score']:.1f}"
        if ja and "score" in severity
        else f" | CVSS 3 base {severity['score']:.1f}"
        if "score" in severity
        else ""
    )
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"{'重要度' if ja else 'Severity'}  {label(finding, ja)}{score}")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(PALETTE[level][2])
    sources = severity.get("sources", [])
    if not sources:
        document.add_paragraph(
            "重要度を判定できるDB情報がありません。参照先で確認してください。"
            if ja
            else "No usable severity metadata was available. Review the advisory references."
        )
    grouped: dict[str, list[str]] = {}
    for source in sources:
        value = (
            (f"CVSS 3 基本値 {source['score']:.1f}" if ja else f"CVSS 3 base {source['score']:.1f}")
            if "score" in source
            else (f"DB評価 {source['value']}" if ja else f"Database label {source['value']}")
            if source["type"] == "database_label"
            else f"{source['type']} ({'未計算' if ja else 'not calculated'})"
        )
        values = grouped.setdefault(source["recordId"], [])
        if value not in values:
            values.append(value)
    for record_id, values in grouped.items():
        paragraph = document.add_paragraph(f"{record_id} | {' / '.join(values)}")
        for run in paragraph.runs:
            run.font.size = Pt(9)
